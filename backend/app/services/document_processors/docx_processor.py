"""
Word (DOCX) 文件處理器 — 使用 python-docx
"""
import io
from typing import List

from app.services.document_processors.base import (
    BaseDocumentProcessor,
    PageText,
    ProcessedDocument,
)


class DocxProcessor(BaseDocumentProcessor):
    """Word DOCX 文件處理器"""

    @property
    def supported_extensions(self) -> List[str]:
        return [".docx"]

    def extract_text(self, file_bytes: bytes, filename: str = "") -> ProcessedDocument:
        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx is required: pip install python-docx")

        doc = Document(io.BytesIO(file_bytes))
        paragraphs: List[str] = []

        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)

        # DOCX 沒有真正的「頁」概念，以段落群組作為頁
        # 每 ~20 段落為一頁（估算）
        pages: List[PageText] = []
        chunk_size = 20
        for i in range(0, len(paragraphs), chunk_size):
            group = paragraphs[i : i + chunk_size]
            pages.append(PageText(
                page_number=i // chunk_size + 1,
                text="\n".join(group),
            ))

        if not pages and paragraphs:
            pages.append(PageText(page_number=1, text="\n".join(paragraphs)))

        return ProcessedDocument(
            text_pages=pages,
            metadata={
                "source": filename,
                "type": "docx",
                "total_paragraphs": len(paragraphs),
            },
        )
